import csv
import io
from typing import Any


COMPANY_FIELD_ALIASES = {
    "name": ["name", "company", "company name", "company_name", "organization"],
    "industry": ["industry", "sector"],
    "sub_sector": ["sub_sector", "sub sector", "subsector", "sub-sector", "segment"],
    "size": ["size", "company size", "company_size"],
    "geography": ["geography", "location", "region", "country", "hq"],
    "client_status": ["client_status", "client status", "status", "relationship"],
    "website": ["website", "url", "web"],
    "notes": ["notes", "description", "comments"],
}

CONTACT_FIELD_ALIASES = {
    "name": ["name", "contact name", "contact_name", "full name", "full_name"],
    "title": ["title", "job title", "job_title", "role", "position"],
    "email": ["email", "email address", "email_address"],
    "company_name": ["company", "company name", "company_name", "organization"],
    "relationship_strength": ["relationship_strength", "relationship strength", "strength", "rating"],
    "notes": ["notes", "comments", "description"],
}


def _resolve_headers(headers: list[str], aliases: dict[str, list[str]]) -> dict[int, str]:
    mapping = {}
    normalized = [h.strip().lower() for h in headers]
    for field, field_aliases in aliases.items():
        for alias in field_aliases:
            if alias.lower() in normalized:
                idx = normalized.index(alias.lower())
                mapping[idx] = field
                break
    return mapping


def parse_csv_content(content: str, field_aliases: dict[str, list[str]]) -> list[dict[str, Any]]:
    reader = csv.reader(io.StringIO(content))
    headers = next(reader, None)
    if not headers:
        return []

    col_map = _resolve_headers(headers, field_aliases)
    if not col_map:
        return []

    rows = []
    for row in reader:
        record = {}
        for idx, field in col_map.items():
            if idx < len(row) and row[idx].strip():
                record[field] = row[idx].strip()
        if record.get("name"):
            rows.append(record)

    return rows


def _parse_pdf_tables(file_content: bytes, aliases: dict[str, list[str]]) -> list[dict[str, Any]]:
    try:
        import pdfplumber
    except ImportError:
        raise ValueError("PDF support requires pdfplumber: pip install pdfplumber")

    rows: list[dict[str, Any]] = []
    with pdfplumber.open(io.BytesIO(file_content)) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables()
            for table in tables:
                if not table or len(table) < 2:
                    continue
                headers = [str(h or "").strip() for h in table[0]]
                col_map = _resolve_headers(headers, aliases)
                if not col_map:
                    continue
                for data_row in table[1:]:
                    record = {}
                    for idx, field in col_map.items():
                        if idx < len(data_row) and data_row[idx] is not None:
                            record[field] = str(data_row[idx]).strip()
                    if record.get("name"):
                        rows.append(record)
    return rows


def _parse_docx_tables(file_content: bytes, aliases: dict[str, list[str]]) -> list[dict[str, Any]]:
    try:
        from docx import Document
    except ImportError:
        raise ValueError("DOCX support requires python-docx: pip install python-docx")

    rows: list[dict[str, Any]] = []
    doc = Document(io.BytesIO(file_content))
    for table in doc.tables:
        if len(table.rows) < 2:
            continue
        headers = [cell.text.strip() for cell in table.rows[0].cells]
        col_map = _resolve_headers(headers, aliases)
        if not col_map:
            continue
        for row in table.rows[1:]:
            cells = [cell.text.strip() for cell in row.cells]
            record = {}
            for idx, field in col_map.items():
                if idx < len(cells) and cells[idx]:
                    record[field] = cells[idx]
            if record.get("name"):
                rows.append(record)
    return rows


def parse_upload(file_content: bytes, filename: str, entity_type: str) -> list[dict[str, Any]]:
    aliases = COMPANY_FIELD_ALIASES if entity_type == "company" else CONTACT_FIELD_ALIASES
    fn_lower = filename.lower()

    if fn_lower.endswith((".xlsx", ".xls")):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(file_content), read_only=True)
            ws = wb.active
            rows_iter = ws.iter_rows(values_only=True)
            headers = [str(h or "") for h in next(rows_iter, [])]
            col_map = _resolve_headers(headers, aliases)
            if not col_map:
                return []

            rows = []
            for row in rows_iter:
                record = {}
                for idx, field in col_map.items():
                    if idx < len(row) and row[idx] is not None:
                        record[field] = str(row[idx]).strip()
                if record.get("name"):
                    rows.append(record)
            return rows
        except ImportError:
            raise ValueError("Excel support requires openpyxl: pip install openpyxl")

    if fn_lower.endswith(".pdf"):
        return _parse_pdf_tables(file_content, aliases)

    if fn_lower.endswith((".docx", ".doc")):
        return _parse_docx_tables(file_content, aliases)

    content = file_content.decode("utf-8-sig")
    return parse_csv_content(content, aliases)
